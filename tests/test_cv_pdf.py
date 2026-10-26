import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi import HTTPException

from jobbot.api import download_cv_pdf
from jobbot.cv.ai import GeneratedText
from jobbot.cv.generator import generate_cv, TEMPLATE_PATH, OUTPUT_DIR
from jobbot.cv.pdf import cv_pdf_bytes, PDFConversionError


class CVPDFTests(unittest.TestCase):
    def test_only_profile_text_changes_in_template(self):
        OUTPUT_DIR.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=OUTPUT_DIR) as directory:
            with patch("jobbot.cv.generator.OUTPUT_DIR", Path(directory)), patch("jobbot.cv.generator.load_profile", return_value={}):
                with patch("jobbot.cv.generator.generate_profile", return_value=GeneratedText("Perfil con C# & .NET < APIs", "test/model")):
                    result = generate_cv("format-test", "Developer")
            original = Document(TEMPLATE_PATH)
            generated = Document(result["file_path"])
            self.assertEqual(
                [p.text.replace("{{ perfil }}", result["profile_text"]) for p in original.paragraphs],
                [p.text for p in generated.paragraphs],
            )
            self.assertEqual(original.styles.element.xml, generated.styles.element.xml)
            self.assertEqual(original.sections[0]._sectPr.xml, generated.sections[0]._sectPr.xml)

    def test_pdf_export_reads_converter_output(self):
        OUTPUT_DIR.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=OUTPUT_DIR) as directory:
            source = Path(directory) / "cv.docx"
            source.write_bytes(b"docx")

            def convert(command, **kwargs):
                destination = Path(command[command.index("--outdir") + 1])
                (destination / "cv.pdf").write_bytes(b"%PDF-test")
                return type("Result", (), {"returncode": 0})()

            with patch("jobbot.cv.pdf.find_libreoffice", return_value="soffice"):
                with patch("jobbot.cv.pdf.subprocess.run", side_effect=convert) as run:
                    self.assertEqual(cv_pdf_bytes(str(source.resolve())), b"%PDF-test")
                    self.assertEqual(run.call_args.args[0][-1], str(source.resolve()))

    def test_download_uses_stored_docx(self):
        OUTPUT_DIR.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=OUTPUT_DIR) as directory:
            source = Path(directory).resolve() / "saved.docx"
            source.write_bytes(b"docx")

            def convert(command, **kwargs):
                destination = Path(command[command.index("--outdir") + 1])
                (destination / "saved.pdf").write_bytes(b"%PDF-test")

            with patch("jobbot.api.repo.get_cv_version_detail", return_value={"file_path": str(source)}):
                with patch("jobbot.api.subprocess.run", side_effect=convert) as run:
                    response = download_cv_pdf("version", db=None)
                    self.assertEqual(run.call_args.args[0][-1], str(source))
                    self.assertEqual(response.media_type, "application/pdf")

    def test_missing_docx_does_not_return_partial_cv(self):
        with patch("jobbot.api.repo.get_cv_version_detail", return_value={"file_path": "generated/missing.docx"}):
            with self.assertRaises(HTTPException) as caught:
                download_cv_pdf("version", db=None)
            self.assertEqual(caught.exception.status_code, 404)

    def test_converter_error_is_reported(self):
        OUTPUT_DIR.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=OUTPUT_DIR) as directory:
            source = Path(directory).resolve() / "saved.docx"
            source.write_bytes(b"docx")
            with patch("jobbot.api.repo.get_cv_version_detail", return_value={"file_path": str(source)}):
                with patch("jobbot.api.subprocess.run", side_effect=FileNotFoundError("Falta soffice")):
                    with self.assertRaises(HTTPException) as caught:
                        download_cv_pdf("version", db=None)
                    self.assertEqual(caught.exception.status_code, 500)


if __name__ == "__main__":
    unittest.main()
