import asyncio
import io
import json
import os
import unittest
from unittest.mock import patch, MagicMock
from pathlib import Path
from tempfile import TemporaryDirectory
from urllib.error import HTTPError, URLError

from cryptography.fernet import Fernet
from docx import Document
from fastapi import HTTPException
from starlette.requests import Request

from jobbot import storage
from jobbot.api import app, get_db
from jobbot.auth import require_user
from jobbot.deployment import validate_production
from jobbot.gmail_oauth import OAuthConfig
from jobbot.gmail_routes import callback_page, protect_local

ENV = {'APP_ENV': 'production', 'AUTH_REQUIRED': 'false', 'SUPABASE_URL': 'https://project.supabase.co',
       'SUPABASE_ANON_KEY': 'public-key', 'SUPABASE_SERVICE_ROLE_KEY': 'server-key',
       'APP_OWNER_ID': 'owner', 'DATABASE_URL': 'postgresql://example/db',
       'FRONTEND_ORIGINS': 'https://jobhunt.example', 'DOCUMENT_STORAGE': 'supabase',
       'GMAIL_OAUTH_REDIRECT_URI': 'https://api.example/api/settings/gmail/callback',
       'MAIL_OAUTH_CLIENT_ID': 'test', 'MAIL_OAUTH_CLIENT_SECRET': 'secret',
       'GMAIL_TOKEN_KEY': Fernet.generate_key().decode(),
       'PROFILE_OBJECT_KEY': 'private/profile.yml', 'CV_TEMPLATE_OBJECT_KEY': 'private/CV_Master.docx'}


def request(path='/api/jobs', headers=(), scheme='https', host='api.example'):
    return Request({'type': 'http', 'path': path, 'scheme': scheme, 'server': (host, 443),
                    'client': ('remote', 123), 'headers': list(headers), 'query_string': b''})


def response(value):
    mock = MagicMock()
    mock.__enter__.return_value.read.return_value = json.dumps(value).encode()
    return mock


class DeploymentTests(unittest.TestCase):
    def setUp(self):
        self.env = patch.dict(os.environ, ENV, clear=True)
        self.env.start()
        self.addCleanup(self.env.stop)

    def test_production_settings_fail_closed(self):
        validate_production()
        del os.environ['APP_OWNER_ID']
        with self.assertRaises(RuntimeError):
            validate_production()

    def test_wildcard_origin_and_local_storage_are_rejected(self):
        os.environ['FRONTEND_ORIGINS'] = '*'
        with self.assertRaises(RuntimeError):
            validate_production()
        os.environ['FRONTEND_ORIGINS'] = ENV['FRONTEND_ORIGINS']
        os.environ['DOCUMENT_STORAGE'] = 'local'
        with self.assertRaises(RuntimeError):
            validate_production()

    def test_production_cannot_disable_auth(self):
        with self.assertRaises(HTTPException) as caught:
            require_user(request())
        self.assertEqual(caught.exception.status_code, 401)

    def test_owner_is_verified_by_auth_server(self):
        with patch('jobbot.auth.urlopen', return_value=response({'id': 'owner'})) as call:
            self.assertEqual(require_user(request(headers=[(b'authorization', b'Bearer token')]))['id'], 'owner')
        self.assertEqual(call.call_args.args[0].full_url, 'https://project.supabase.co/auth/v1/user')

    def test_other_and_anonymous_accounts_are_forbidden(self):
        for user in ({'id': 'someone-else'}, {'id': 'owner', 'is_anonymous': True}):
            with patch('jobbot.auth.urlopen', return_value=response(user)), self.assertRaises(HTTPException) as caught:
                require_user(request(headers=[(b'authorization', b'Bearer token')]))
            self.assertEqual(caught.exception.status_code, 403)

    def test_invalid_token_and_unavailable_auth_fail_closed(self):
        for error, expected in ((HTTPError('https://auth', 401, 'bad', {}, None), 401), (URLError('network'), 503)):
            with patch('jobbot.auth.urlopen', side_effect=error), self.assertRaises(HTTPException) as caught:
                require_user(request(headers=[(b'authorization', b'Bearer token')]))
            self.assertEqual(caught.exception.status_code, expected)

    def test_every_data_route_is_protected_before_database_access(self):
        async def call(path, method):
            messages = []
            async def receive():
                return {'type': 'http.request', 'body': b'{}', 'more_body': False}
            async def send(message):
                messages.append(message)
            scope = dict(request(path).scope, method=method, http_version='1.1', root_path='', raw_path=path.encode())
            await app(scope, receive, send)
            return next(m['status'] for m in messages if m['type'] == 'http.response.start')
        def forbidden_db():
            raise AssertionError('Una petición anónima no debe abrir la base de datos')
        app.dependency_overrides[get_db] = forbidden_db
        self.addCleanup(app.dependency_overrides.clear)
        for path, method in (('/api/jobs', 'GET'), ('/api/jobs/page', 'GET'), ('/api/stats', 'GET'),
                             ('/api/cvs', 'GET'), ('/api/cvs/x/pdf', 'GET'), ('/api/jobs/x/cv', 'POST'),
                             ('/api/cvs/x/cover-letter', 'POST'), ('/api/settings/gmail/connect', 'POST'),
                             ('/api/collection/email', 'GET')):
            self.assertEqual(asyncio.run(call(path, method)), 401, path)
        self.assertEqual(asyncio.run(call('/healthz', 'GET')), 200)

    def test_remote_oauth_uses_https_and_secure_cookies(self):
        self.assertTrue(OAuthConfig.load().redirect_uri.startswith('https://'))
        protect_local(request('/api/settings/gmail/callback'))
        with self.assertRaises(HTTPException):
            protect_local(request('/api/settings/gmail/callback', scheme='http'))
        with self.assertRaises(HTTPException):
            protect_local(request('/api/settings/gmail/callback', host='attacker.example'))
        cookies = callback_page('https://jobhunt.example', False, 'Error').headers.getlist('set-cookie')
        self.assertTrue(all('Secure' in c and 'HttpOnly' in c and 'SameSite=lax' in c for c in cookies))

    def test_wrong_origin_cannot_mutate_gmail(self):
        with self.assertRaises(HTTPException):
            protect_local(request(headers=[(b'origin', b'https://attacker.example'), (b'content-type', b'application/json')]), mutation=True)

    def test_gmail_status_can_report_missing_oauth_credentials(self):
        del os.environ['MAIL_OAUTH_CLIENT_ID']
        # Host/origin checks must not require an already complete Google setup.
        protect_local(request('/api/settings/gmail'))

    def test_public_bucket_is_rejected(self):
        with patch('jobbot.storage.urlopen', return_value=response({'public': True})), self.assertRaises(storage.StorageError):
            storage.verify_private_bucket()
        with patch('jobbot.storage.urlopen', return_value=response({'public': False})):
            storage.verify_private_bucket()

    def test_storage_references_cannot_escape_the_private_bucket(self):
        for value in ('storage://other/cv.docx', 'storage://jobhunt-private/../secret', 'file:///etc/passwd'):
            with self.assertRaises(storage.StorageError):
                storage.read(value)

    def test_remote_pdf_uses_private_docx_bytes(self):
        from jobbot.cv.pdf import cv_pdf_bytes
        with patch('jobbot.storage.read', return_value=b'docx'), patch('jobbot.cv.pdf.convert_docx', return_value=b'%PDF-test') as convert:
            self.assertEqual(cv_pdf_bytes('storage://jobhunt-private/cvs/a.docx'), b'%PDF-test')
            convert.assert_called_once_with(b'docx')

    def test_cv_uploads_rendered_docx_and_records_remote_reference(self):
        from jobbot.cv.generator import generate_cv
        from jobbot.cv.ai import GeneratedText
        doc = Document()
        doc.add_paragraph('{{ perfil }}')
        template = io.BytesIO()
        doc.save(template)
        ref = 'storage://jobhunt-private/cvs/result.docx'
        with patch('jobbot.storage.read', return_value=template.getvalue()), patch('jobbot.storage.write', return_value=ref) as upload, \
             patch('jobbot.cv.generator.load_profile', return_value={'skills': ['Python']}), \
             patch('jobbot.cv.generator.generate_profile', return_value=GeneratedText('Perfil real', 'test/model')):
            result = generate_cv('job', 'Developer')
        self.assertEqual(result['file_path'], ref)
        rendered = Document(io.BytesIO(upload.call_args.args[1]))
        self.assertEqual(rendered.paragraphs[0].text, 'Perfil real')

    def test_storage_embedded_404_is_reported_as_missing(self):
        body = io.BytesIO(b'{"statusCode":"404","error":"not_found"}')
        with patch('jobbot.storage.urlopen', side_effect=HTTPError('https://storage', 400, 'missing', {}, body)), self.assertRaises(FileNotFoundError):
            storage.read('storage://jobhunt-private/cvs/missing.docx')

    def test_document_migration_is_preview_first_and_repeatable(self):
        from jobbot.db import Database
        from jobbot.models import JobPosting
        from jobbot import repo
        from jobbot.cv.generator import OUTPUT_DIR
        from jobbot.cv.profile import PROJECT_ROOT
        from jobbot.migrate_documents import migrate
        db = Database('sqlite:///:memory:')
        db.init_schema()
        self.addCleanup(db.close)
        repo.insert_job_if_new(db, JobPosting(source='test', title='Developer', url='https://example.com/job'))
        job_id = db.query('select id from jobs')[0]['id']
        OUTPUT_DIR.mkdir(exist_ok=True)
        with TemporaryDirectory(dir=OUTPUT_DIR) as directory:
            path = Path(directory).resolve() / 'cv.docx'
            path.write_bytes(b'docx')
            original = path.relative_to(PROJECT_ROOT).as_posix()
            version = repo.save_cv_version(db, job_id, 'ai', 'Perfil', original, 'ai', 'test/model')
            with patch('jobbot.storage.write') as upload:
                self.assertEqual(migrate(db)['pending'], 1)
                upload.assert_not_called()
            ref = 'storage://jobhunt-private/cvs/imported/cv.docx'
            with patch('jobbot.storage.read', side_effect=FileNotFoundError), patch('jobbot.storage.write', return_value=ref):
                self.assertEqual(migrate(db, True)['migrated'], 1)
            self.assertEqual(db.query('select file_path from cv_versions where id = ?', (version['id'],))[0]['file_path'], ref)
            self.assertEqual(migrate(db, True)['remote'], 1)
            self.assertTrue(path.is_file())

    def test_failed_upload_does_not_rewrite_existing_cv_reference(self):
        from jobbot.db import Database
        from jobbot.models import JobPosting
        from jobbot import repo
        from jobbot.cv.generator import OUTPUT_DIR
        from jobbot.cv.profile import PROJECT_ROOT
        from jobbot.migrate_documents import migrate
        db = Database('sqlite:///:memory:')
        db.init_schema()
        self.addCleanup(db.close)
        repo.insert_job_if_new(db, JobPosting(source='test', title='Developer', url='https://example.com/job'))
        job_id = db.query('select id from jobs')[0]['id']
        OUTPUT_DIR.mkdir(exist_ok=True)
        with TemporaryDirectory(dir=OUTPUT_DIR) as directory:
            path = Path(directory).resolve() / 'cv.docx'
            path.write_bytes(b'docx')
            original = path.relative_to(PROJECT_ROOT).as_posix()
            repo.save_cv_version(db, job_id, 'ai', 'Perfil', original, 'ai', 'test/model')
            with patch('jobbot.storage.read', side_effect=FileNotFoundError), patch('jobbot.storage.write', side_effect=storage.StorageError('upload failed')):
                with self.assertRaises(storage.StorageError):
                    migrate(db, True)
            self.assertEqual(db.query('select file_path from cv_versions')[0]['file_path'], original)
